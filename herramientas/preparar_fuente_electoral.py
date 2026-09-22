#!/usr/bin/env python3
from __future__ import annotations
import argparse,csv,hashlib,json,re,shutil,unicodedata
from datetime import datetime,timezone
from pathlib import Path
import yaml
from herramientas.comprobar_fuente_electoral_oficial import check_declaration,load_declaration


def _clean_header(value: object) -> str:
    text=unicodedata.normalize("NFKD",str(value or "")).encode("ascii","ignore").decode("ascii").strip().lower()
    return re.sub(r"[^a-z0-9]+","_",text).strip("_")


def _as_int(value: object) -> int:
    if value in (None,""): return 0
    if isinstance(value,(int,float)): return int(value)
    raw=str(value).strip().replace(".","").replace(",","")
    return int(float(raw or 0))


def _normalise_code(value: object,width:int) -> str:
    if value in (None,""): return ""
    raw=str(value).strip()
    try:
        raw=str(int(float(raw)))
    except Exception:
        raw=re.sub(r"\D","",raw)
    return raw.zfill(width)


def _raw_code(value: object) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def transform_gipeyop_polling_xlsx(source:Path,out_dir:Path,declaration:dict,source_decl:dict)->dict:
    try:
        import openpyxl
    except Exception as exc:
        raise RuntimeError("La transformación XLSX electoral requiere openpyxl") from exc
    sheet_name=str((source_decl.get("transform") or {}).get("sheet") or "MESAS")
    wb=openpyxl.load_workbook(source,read_only=True,data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"XLSX sin hoja {sheet_name!r}: {source}")
    ws=wb[sheet_name]
    rows=ws.iter_rows(values_only=True)
    try: headers_raw=next(rows)
    except StopIteration: raise ValueError("XLSX electoral vacío")
    headers=[_clean_header(x) for x in headers_raw]
    aliases={
        "year":{"year","anyo","ano"}, "province":{"cod_prov","codigo_provincia","provincia"},
        "municipality":{"cod_mun","codigo_municipio","municipio"}, "district":{"distrito","codigo_distrito"},
        "section":{"seccion","codigo_seccion"}, "polling":{"mesa","codigo_mesa"},
        "total_votes":{"votos","votos_totales"}, "blank":{"blancos","votos_blancos"}, "null":{"nulos","votos_nulos"},
    }
    positions={}
    for key,names in aliases.items():
        for idx,name in enumerate(headers):
            if name in names:
                positions[key]=idx; break
    required={"province","municipality","district","section","total_votes","blank","null"}
    missing=sorted(required-set(positions))
    if missing: raise ValueError(f"XLSX electoral sin columnas estructurales: {missing}; cabecera={headers}")
    party_start=max(positions["total_votes"],positions["blank"],positions["null"])+1
    party_headers=[]
    for idx in range(party_start,len(headers_raw)):
        raw=str(headers_raw[idx] or "").strip()
        if raw: party_headers.append((idx,raw))
    if not party_headers: raise ValueError("XLSX electoral sin columnas de candidaturas")
    expected_year=str((source_decl.get("transform") or {}).get("election_year") or str(declaration.get("election_date") or "")[:4])
    out_dir.mkdir(parents=True,exist_ok=True)
    normalized=out_dir/"resultados_electorales_normalizados.csv"
    parties=set(); records=0; sections=set()
    with normalized.open("w",encoding="utf-8",newline="") as fh:
        writer=csv.DictWriter(fh,fieldnames=["CUSEC_KEY","party","votes"],delimiter=";")
        writer.writeheader()
        for values in rows:
            if not values: continue
            if "year" in positions and expected_year:
                y=str(values[positions["year"]] or "").strip()
                try: y=str(int(float(y)))
                except Exception: pass
                if y and y!=expected_year: continue
            prov=_normalise_code(values[positions["province"]],2)
            mun=_normalise_code(values[positions["municipality"]],3)
            dist=_normalise_code(values[positions["district"]],2)
            sec_raw=_raw_code(values[positions["section"]])
            sec=_normalise_code(values[positions["section"]],3)
            if not (prov and mun and dist and sec_raw): continue
            # CERA codificada como municipio 991/992/993 queda fuera de la cartografía,
            # pero no se mezcla con secciones ordinarias.
            if mun in {"991","992","993","999"}: continue
            # Códigos especiales de sección (0, negativos o alfanuméricos) no son CUSEC.
            # Se preservan como identificadores explícitos para que reconciliación los
            # contabilice como result_only; nunca deben colisionar con una sección real.
            special_section = (
                not re.fullmatch(r"\d+", sec_raw)
                or int(sec_raw) <= 0
            )
            cusec=(f"SPECIAL:{prov}:{mun}:{dist}:{sec_raw}" if special_section else f"{prov}{mun}{dist}{sec}")
            sections.add(cusec)
            for idx,party in party_headers:
                votes=_as_int(values[idx] if idx < len(values) else 0)
                parties.add(party)
                writer.writerow({"CUSEC_KEY":cusec,"party":party,"votes":votes})
                records+=1
    wb.close()
    if not records or not sections: raise ValueError("La transformación electoral no produjo registros censales")
    dictionary=out_dir/"party_dictionary.json"
    dictionary_payload={
        "schema_family":"ddd-party-dictionary","schema_version":"1.0.0","unknown_party_policy":"reject",
        "parties":[{"canonical_id":p,"display_name":p,"aliases":[],"classification":""} for p in sorted(parties)],
    }
    dictionary.write_text(json.dumps(dictionary_payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    contract=out_dir/"election_contract.json"
    contract_payload={
        "schema_family":"ddd-election","schema_version":"1.0.0",
        "election_id":str(declaration["election_id"]),"territory_id":str(declaration["territory_id"]),
        "title":str(declaration.get("title") or declaration["election_id"]),
        "election_date":str(declaration["election_date"]),"input_mode":"verifiable_file","boundary_independence":True,
        "sources":[{
            "path":normalized.as_posix(),"sha256":sha(normalized),
            "publisher":str(source_decl.get("publisher") or ""),"source_url":str(source_decl.get("url") or ""),
            "retrieved_at":datetime.now(timezone.utc).date().isoformat(),
            "adapter":{"kind":"long_csv","separator":";","section_field":"CUSEC_KEY","party_field":"party","votes_field":"votes"},
        }],
        "party_dictionary":{"path":dictionary.as_posix(),"sha256":sha(dictionary)},
        "reconciliation":declaration.get("reconciliation") or {"policy":"fail_unless_declared","allowed_result_only_sections":[],"allowed_map_only_sections":[]},
    }
    contract.write_text(json.dumps(contract_payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    return {"source":normalized,"contract":contract,"dictionary":dictionary,"records":records,"sections":len(sections),"parties":len(parties)}


def transform_minsait_polling_long_csv(source:Path,out_dir:Path,declaration:dict,source_decl:dict)->dict:
    transform=source_decl.get("transform") or {}
    fields,rows=_read_delimited(source)
    clean_to_raw={_clean_header(name):name for name in fields}

    def resolve_field(config_key:str,defaults:list[str],required:bool=True)->str|None:
        configured=transform.get(config_key)
        candidates=[]
        if isinstance(configured,list):
            candidates.extend(str(x) for x in configured)
        elif configured:
            candidates.append(str(configured))
        candidates.extend(defaults)
        for candidate in candidates:
            raw=clean_to_raw.get(_clean_header(candidate))
            if raw:
                return raw
        if required:
            raise ValueError(f"CSV Minsait sin campo {config_key}; cabecera={fields}")
        return None

    province_field=resolve_field("province_field",["codigo_provincia","provincia"])
    municipality_field=resolve_field("municipality_field",["codigo_municipio","municipio"])
    district_field=resolve_field("district_field",["codigo_distrito","distrito"])
    section_field=resolve_field("section_field",["codigo_seccion","seccion"])
    polling_field=resolve_field("polling_station_field",["codigo_mesa","mesa"],required=False)
    votes_field=resolve_field("votes_field",["votos","votes"])
    party_field=resolve_field("party_fields",["recode","partido","siglas","denominacion"])
    ccaa_field=resolve_field("autonomous_community_field",["codigo_ccaa"],required=False)

    expected_ccaa=str(transform.get("autonomous_community_code") or "").strip()
    province_width=int(transform.get("province_width",2))
    municipality_width=int(transform.get("municipality_width",3))
    district_width=int(transform.get("district_width",2))
    section_width=int(transform.get("section_width",3))

    aggregated={}
    sections=set()
    polling_stations=set()
    skipped=0
    for row in rows:
        if ccaa_field and expected_ccaa:
            raw_ccaa=_normalise_code(row.get(ccaa_field),len(expected_ccaa))
            if raw_ccaa and raw_ccaa!=expected_ccaa.zfill(len(expected_ccaa)):
                continue
        prov=_normalise_code(row.get(province_field),province_width)
        mun=_normalise_code(row.get(municipality_field),municipality_width)
        dist=_normalise_code(row.get(district_field),district_width)
        sec_raw=_raw_code(row.get(section_field))
        sec=_normalise_code(row.get(section_field),section_width)
        party=str(row.get(party_field) or "").strip()
        if not (prov and mun and dist and sec_raw and party):
            skipped+=1
            continue
        if mun in {"991","992","993","999"}:
            skipped+=1
            continue
        special_section=(not re.fullmatch(r"\d+",sec_raw) or int(sec_raw)<=0 or len(re.sub(r"\D","",sec_raw))>section_width)
        cusec=(f"SPECIAL:{prov}:{mun}:{dist}:{sec_raw}" if special_section else f"{prov}{mun}{dist}{sec}")
        votes=_as_int(row.get(votes_field))
        key=(cusec,party)
        aggregated[key]=aggregated.get(key,0)+votes
        sections.add(cusec)
        if polling_field:
            mesa=str(row.get(polling_field) or "").strip()
            if mesa:
                polling_stations.add((cusec,mesa))

    if not aggregated or not sections:
        raise ValueError("La transformación Minsait no produjo votos por sección")
    out_dir.mkdir(parents=True,exist_ok=True)
    normalized=out_dir/"resultados_electorales_normalizados.csv"
    with normalized.open("w",encoding="utf-8",newline="") as fh:
        writer=csv.DictWriter(fh,fieldnames=["CUSEC_KEY","party","votes"],delimiter=";")
        writer.writeheader()
        for (cusec,party),votes in sorted(aggregated.items()):
            writer.writerow({"CUSEC_KEY":cusec,"party":party,"votes":votes})

    parties=sorted({party for _,party in aggregated})
    dictionary=out_dir/"party_dictionary.json"
    dictionary_payload={
        "schema_family":"ddd-party-dictionary","schema_version":"1.0.0","unknown_party_policy":"reject",
        "parties":[{"canonical_id":p,"display_name":p,"aliases":[],"classification":""} for p in parties],
    }
    dictionary.write_text(json.dumps(dictionary_payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    contract=out_dir/"election_contract.json"
    contract_payload={
        "schema_family":"ddd-election","schema_version":"1.0.0",
        "election_id":str(declaration["election_id"]),"territory_id":str(declaration["territory_id"]),
        "title":str(declaration.get("title") or declaration["election_id"]),
        "election_date":str(declaration["election_date"]),"input_mode":"verifiable_file","boundary_independence":True,
        "sources":[{
            "path":normalized.as_posix(),"sha256":sha(normalized),
            "publisher":str(source_decl.get("publisher") or ""),"source_url":str(source_decl.get("url") or ""),
            "retrieved_at":datetime.now(timezone.utc).date().isoformat(),
            "adapter":{"kind":"long_csv","separator":";","section_field":"CUSEC_KEY","party_field":"party","votes_field":"votes"},
        }],
        "party_dictionary":{"path":dictionary.as_posix(),"sha256":sha(dictionary)},
        "reconciliation":declaration.get("reconciliation") or {"policy":"fail_unless_declared","allowed_result_only_sections":[],"allowed_map_only_sections":[]},
    }
    contract.write_text(json.dumps(contract_payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    return {
        "source":normalized,"contract":contract,"dictionary":dictionary,
        "records":len(aggregated),"sections":len(sections),"parties":len(parties),
        "polling_stations":len(polling_stations),"raw_rows":len(rows),"skipped_rows":skipped,
    }


def _transform_selected_source(src:Path,tmp:Path,declaration:dict,source_decl:dict)->dict|None:
    transform=source_decl.get("transform") or {}
    kind=str(transform.get("kind") or "")
    if not kind: return None
    if kind=="gipeyop_polling_xlsx":
        return transform_gipeyop_polling_xlsx(src,tmp/"normalized",declaration,source_decl)
    if kind=="minsait_polling_long_csv":
        return transform_minsait_polling_long_csv(src,tmp/"normalized",declaration,source_decl)
    raise ValueError(f"Transformación electoral no soportada: {kind}")

def sha(path:Path)->str:
    h=hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda:f.read(1024*1024),b""): h.update(chunk)
    return h.hexdigest()

def count_records(path:Path)->tuple[int,str]:
    try:
        if path.suffix.lower() in {".csv",".tsv",".txt"}:
            text=path.read_text(encoding="utf-8-sig",errors="strict")
            delim="\t" if path.suffix.lower()==".tsv" else ("," if text.splitlines() and text.splitlines()[0].count(",")>=text.splitlines()[0].count(";") else ";")
            rows=list(csv.reader(text.splitlines(),delimiter=delim))
            return max(0,len(rows)-1),"delimited_rows_excluding_header"
        if path.suffix.lower()==".json":
            obj=json.loads(path.read_text(encoding="utf-8"))
            if isinstance(obj,list): return len(obj),"json_top_level_list"
            cur=obj
            for key in ("mapa","zonas"):
                if isinstance(cur,dict) and key in cur: cur=cur[key]
            if isinstance(cur,list): return len(cur),"json_mapa_zonas"
            if isinstance(obj,dict) and isinstance(obj.get("features"),list): return len(obj["features"]),"geojson_features"
    except Exception:
        pass
    return 0,"not_countable"

def _read_delimited(path:Path)->tuple[list[str],list[dict]]:
    raw=path.read_bytes()
    text=None
    for encoding in ("utf-8-sig","utf-8","latin-1"):
        try:
            text=raw.decode(encoding); break
        except UnicodeDecodeError:
            continue
    if text is None: raise ValueError(f"No se puede decodificar {path}")
    lines=text.splitlines()
    if not lines: raise ValueError(f"Fuente vacía: {path}")
    delimiter=max((";",",","\t"),key=lambda d:lines[0].count(d))
    reader=csv.DictReader(lines,delimiter=delimiter)
    fields=[str(x or "").strip() for x in (reader.fieldnames or [])]
    if not fields: raise ValueError(f"CSV sin cabecera: {path}")
    rows=[]
    for row in reader:
        rows.append({fields[i]: row.get(reader.fieldnames[i],"") for i in range(len(fields))})
    return fields,rows

def merge_delimited_sources(paths:list[Path],out:Path)->dict:
    all_fields=[]; all_rows=[]
    for path in paths:
        fields,rows=_read_delimited(path)
        for field in fields:
            if field not in all_fields: all_fields.append(field)
        all_rows.extend(rows)
    if not all_rows: raise ValueError("Las fuentes electorales no contienen registros")
    out.parent.mkdir(parents=True,exist_ok=True)
    with out.open("w",encoding="utf-8",newline="") as f:
        writer=csv.DictWriter(f,fieldnames=all_fields,delimiter=";",extrasaction="ignore")
        writer.writeheader()
        for row in all_rows: writer.writerow({field:row.get(field,"") for field in all_fields})
    return {"records":len(all_rows),"columns":all_fields}

def _write_package(out:Path,decision:str,territory_id:str,edition:str,source:Path|None,meta:dict,extra:dict|None=None)->dict:
    if out.exists(): shutil.rmtree(out)
    out.mkdir(parents=True)
    selected=None
    if source is not None:
        data_dir=out/"data"; data_dir.mkdir()
        target=data_dir/source.name; shutil.copy2(source,target)
        records,method=count_records(target)
        selected={"path":target.relative_to(out).as_posix(),"sha256":sha(target),"bytes":target.stat().st_size,"records":records,"record_count_method":method,**meta}
    manifest={"schema":"ddd-electoral-package/1.0","decision":decision,"territory_id":territory_id,"edition":str(edition),"frozen_at":datetime.now(timezone.utc).isoformat(),"selected_source":selected}
    if extra: manifest.update(extra)
    (out/"manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    (out/"decision.json").write_text(json.dumps({"decision":decision,"territory_id":territory_id,"edition":str(edition),"selected_source":selected},ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    return manifest

def _expected_election_identity(root:Path,cfg:dict,declaration:Path|None)->tuple[str|None,str|None]:
    decl=declaration
    if decl is None:
        raw=(cfg.get("meta") or {}).get("electoral_sources_declaration")
        if raw:
            decl=root/str(raw)
    if decl and decl.is_file():
        data=load_declaration(decl)
        return (
            str(data.get("election_id") or "") or None,
            str(data.get("election_date") or "") or None,
        )
    m07=(cfg.get("modulos") or {}).get("modulo_07_agregar_resultados_electorales") or {}
    contract_raw=m07.get("election_contract")
    if contract_raw:
        contract_path=root/str(contract_raw)
        if contract_path.is_file():
            data=json.loads(contract_path.read_text(encoding="utf-8"))
            return (
                str(data.get("election_id") or "") or None,
                str(data.get("election_date") or "") or None,
            )
    return None,None


def validate_previous(package:Path,territory_id:str,edition:str,expected_election_id:str|None=None,expected_election_date:str|None=None)->dict|None:
    try:
        m=json.loads((package/"manifest.json").read_text(encoding="utf-8"))
        if m.get("schema")!="ddd-electoral-package/1.0" or m.get("decision") not in {"REUSE","ACQUIRE"}: return None
        if m.get("territory_id")!=territory_id or str(m.get("edition"))!=str(edition): return None
        if expected_election_id is not None and str(m.get("election_id") or "")!=str(expected_election_id): return None
        if expected_election_date is not None and str(m.get("election_date") or "")!=str(expected_election_date): return None
        s=m.get("selected_source") or {}; p=package/str(s.get("path") or "")
        if not p.is_file() or sha(p)!=str(s.get("sha256") or "") or p.stat().st_size!=int(s.get("bytes") or -1): return None
        records,_=count_records(p)
        if "records" in s and int(s.get("records") or 0)!=records: return None
        return m
    except Exception: return None

def prepare(*,territory_id:str,edition:str,package_out:Path,root:Path,params:Path|None=None,declaration:Path|None=None,previous:Path|None=None,previous_run_id:str|None=None,previous_artifact_name:str|None=None)->dict:
    root=root.resolve()
    cfg={}
    if params and params.is_file(): cfg=yaml.safe_load(params.read_text(encoding="utf-8")) or {}
    expected_election_id,expected_election_date=_expected_election_identity(root,cfg,declaration)
    if previous and previous.is_dir():
        m=validate_previous(
            previous,territory_id,edition,
            expected_election_id=expected_election_id,
            expected_election_date=expected_election_date,
        )
        if m and previous_run_id and previous_artifact_name:
            source=previous/m["selected_source"]["path"]
            meta={k:v for k,v in m["selected_source"].items() if k not in {"path","sha256","bytes","records","record_count_method"}}
            manifest=_write_package(package_out,"REUSE",territory_id,edition,source,meta,{
                "reuse_provenance":{
                    "run_id":str(previous_run_id),
                    "artifact_name":str(previous_artifact_name),
                },
                "election_id":m.get("election_id"),
                "election_date":m.get("election_date"),
            })
            embedded=m.get("embedded_contract") or {}
            if embedded:
                contract_src=previous/str(embedded.get("election_contract") or "")
                dictionary_src=previous/str(embedded.get("party_dictionary") or "")
                if not contract_src.is_file() or not dictionary_src.is_file():
                    raise ValueError("Paquete electoral reutilizable perdió su contrato embebido")
                contract_dir=package_out/"contract"; contract_dir.mkdir(exist_ok=True)
                shutil.copy2(contract_src,contract_dir/"election_contract.json")
                shutil.copy2(dictionary_src,contract_dir/"party_dictionary.json")
                manifest["embedded_contract"]={
                    "election_contract":"contract/election_contract.json",
                    "party_dictionary":"contract/party_dictionary.json",
                    "contract_sha256":sha(contract_dir/"election_contract.json"),
                    "party_dictionary_sha256":sha(contract_dir/"party_dictionary.json"),
                }
                (package_out/"manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
            return manifest
    m07=(cfg.get("modulos") or {}).get("modulo_07_agregar_resultados_electorales") or {}
    contract_raw=m07.get("election_contract")
    if contract_raw:
        contract_path=root/str(contract_raw)
        if contract_path.is_file():
            contract=json.loads(contract_path.read_text(encoding="utf-8"))
            contract_tid=str(contract.get("territory_id") or "")
            contract_election_id=str(contract.get("election_id") or "")
            contract_election_date=str(contract.get("election_date") or "")
            identity_matches=(
                contract_tid == territory_id
                and bool(contract_election_id)
                and bool(contract_election_date)
                and (expected_election_id is None or contract_election_id == expected_election_id)
                and (expected_election_date is None or contract_election_date == expected_election_date)
            )
            sources=contract.get("sources") or []
            if identity_matches and len(sources)==1:
                s=sources[0]; src=root/str(s.get("path") or "")
                expected=str(s.get("sha256") or "").lower()
                if src.is_file() and expected and sha(src).lower()==expected:
                    meta={"origin_url":s.get("source_url"),"publisher":s.get("publisher"),"acquired_at":s.get("retrieved_at"),"source_mode":"existing_contract","contract":str(contract_raw)}
                    return _write_package(
                        package_out,"REUSE",territory_id,edition,src,meta,{
                            "election_id":contract_election_id,
                            "election_date":contract_election_date,
                        },
                    )
    decl=declaration
    if decl is None:
        raw=(cfg.get("meta") or {}).get("electoral_sources_declaration")
        if raw: decl=root/str(raw)
    if decl and decl.is_file():
        tmp=package_out.parent/(package_out.name+"-acquire")
        if tmp.exists(): shutil.rmtree(tmp)
        d=load_declaration(decl); result=check_declaration(d,tmp)
        if result.get("decision")=="READY":
            selected_sources=result.get("selected_sources") or []
            election_extra={"checker_decision":result,"election_id":d.get("election_id"),"election_date":d.get("election_date")}
            if len(selected_sources)>1:
                source_paths=[tmp/str(sel["artifact_path"]) for sel in selected_sources]
                merged=tmp/"resultados_electorales_vigentes.csv"
                merge_info=merge_delimited_sources(source_paths,merged)
                provenance=[{
                    "id":sel.get("id"),"url":sel.get("url"),"publisher":sel.get("publisher"),
                    "sha256":sel.get("sha256"),"bytes":sel.get("bytes")
                } for sel in selected_sources]
                meta={
                    "publisher":"; ".join(sorted({str(sel.get("publisher") or "") for sel in selected_sources if sel.get("publisher")})),
                    "acquired_at":datetime.now(timezone.utc).isoformat(),
                    "source_mode":"official_acquisition_multisource",
                    "source_count":len(selected_sources),
                    "sources":provenance,
                    "merge":merge_info,
                    "declaration":str(decl),
                }
                manifest=_write_package(package_out,"ACQUIRE",territory_id,edition,merged,meta,election_extra)
                raw_dir=package_out/"raw"; raw_dir.mkdir(exist_ok=True)
                for sel,src in zip(selected_sources,source_paths):
                    raw_target=raw_dir/src.name
                    shutil.copy2(src,raw_target)
                manifest["raw_sources"]=[{
                    "id":sel.get("id"),
                    "path":f"raw/{(tmp/str(sel['artifact_path'])).name}",
                    "sha256":sel.get("sha256"),
                    "bytes":sel.get("bytes"),
                    "url":sel.get("url"),
                    "publisher":sel.get("publisher"),
                } for sel in selected_sources]
                (package_out/"manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
                shutil.rmtree(tmp,ignore_errors=True); return manifest
            sel=result.get("selected_source") or (selected_sources[0] if selected_sources else None)
            if not sel:
                raise ValueError("READY sin fuente electoral seleccionada")
            src=tmp/str(sel["artifact_path"])
            declared_source=next((s for s in (d.get("sources") or []) if str(s.get("id") or "")==str(sel.get("id") or "")),{})
            transformed=_transform_selected_source(src,tmp,d,declared_source)
            selected_src=transformed["source"] if transformed else src
            source_mode="verified_mirror_transformed" if transformed and str(declared_source.get("source_class") or "")=="verified_mirror" else ("official_acquisition_transformed" if transformed else "official_acquisition")
            meta={"origin_url":sel.get("url"),"publisher":sel.get("publisher"),"acquired_at":datetime.now(timezone.utc).isoformat(),"source_mode":source_mode,"declaration":str(decl),"source_class":declared_source.get("source_class","official")}
            if transformed:
                meta["transform"]={
                    "kind":str((declared_source.get("transform") or {}).get("kind") or ""),
                    "records":transformed["records"],"sections":transformed["sections"],"parties":transformed["parties"],
                }
                for key in ("polling_stations","raw_rows","skipped_rows"):
                    if key in transformed:
                        meta["transform"][key]=transformed[key]
            manifest=_write_package(package_out,"ACQUIRE",territory_id,edition,selected_src,meta,election_extra)
            if transformed:
                contract_dir=package_out/"contract"; contract_dir.mkdir(exist_ok=True)
                shutil.copy2(transformed["contract"],contract_dir/"election_contract.json")
                shutil.copy2(transformed["dictionary"],contract_dir/"party_dictionary.json")
                embedded={
                    "election_contract":"contract/election_contract.json",
                    "party_dictionary":"contract/party_dictionary.json",
                    "contract_sha256":sha(contract_dir/"election_contract.json"),
                    "party_dictionary_sha256":sha(contract_dir/"party_dictionary.json"),
                }
                manifest["embedded_contract"]=embedded
                # Reescribir paths internos del contrato a rutas portables dentro del paquete.
                ep=json.loads((contract_dir/"election_contract.json").read_text(encoding="utf-8"))
                ep["sources"][0]["path"]="data/"+selected_src.name
                ep["party_dictionary"]["path"]="contract/party_dictionary.json"
                (contract_dir/"election_contract.json").write_text(json.dumps(ep,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
                embedded["contract_sha256"]=sha(contract_dir/"election_contract.json")
                manifest["embedded_contract"]=embedded
                (package_out/"manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
            shutil.rmtree(tmp,ignore_errors=True); return manifest
        manifest=_write_package(package_out,"BLOCK",territory_id,edition,None,{},{"reason":"No existe fuente electoral oficial reutilizable o adquirible","checker_decision":result})
        shutil.copytree(tmp,package_out/"checker",dirs_exist_ok=True)
        shutil.rmtree(tmp,ignore_errors=True)
        return manifest
    return _write_package(package_out,"BLOCK",territory_id,edition,None,{},{"reason":"Sin paquete reutilizable, contrato materializado ni declaración electoral oficial"})

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--territory-id",required=True); ap.add_argument("--edition",required=True)
    ap.add_argument("--package-out",type=Path,required=True); ap.add_argument("--root-dir",type=Path,default=Path("."))
    ap.add_argument("--params",type=Path); ap.add_argument("--declaration",type=Path); ap.add_argument("--previous-package",type=Path)
    ap.add_argument("--previous-run-id"); ap.add_argument("--previous-artifact-name")
    a=ap.parse_args()
    m=prepare(territory_id=a.territory_id,edition=a.edition,package_out=a.package_out,root=a.root_dir,params=a.params,declaration=a.declaration,previous=a.previous_package,previous_run_id=a.previous_run_id,previous_artifact_name=a.previous_artifact_name)
    print(json.dumps(m,ensure_ascii=False))
    if m["decision"]=="BLOCK": raise SystemExit(2)
if __name__=="__main__": main()
